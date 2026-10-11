import os
import openpyxl
from animals import create_animal
from staff import Zookeeper, Vet
from enclosures import Enclosure

# Keep the database next to this file, no matter which folder the program is run from
BASE_DIR = os.path.dirname(os.path.abspath(__file__))

DEFAULT_SALARIES = {"ZooKeeper": 45000, "Veterinarian": 75000}
STAFF_CLASSES = {"ZooKeeper": Zookeeper, "Veterinarian": Vet}


class ZooManager:
    def __init__(self):
        self.animals = []
        self.staff = []
        self.enclosures = []
        self.db_file = os.path.join(BASE_DIR, "zoo_database.xlsx")

        # Automatically load data when the manager starts
        self.load_data()

    # ---------- Lookup helpers ----------

    def find_animal(self, animal_id):
        for animal in self.animals:
            if animal.id == str(animal_id).strip():
                return animal
        return None

    def find_enclosure(self, name):
        for enclosure in self.enclosures:
            if enclosure.name.lower() == str(name).strip().lower():
                return enclosure
        return None

    def find_staff(self, employee_id):
        for member in self.staff:
            if member.employee_id == str(employee_id).strip():
                return member
        return None

    def find_enclosure_of(self, animal):
        for enclosure in self.enclosures:
            if animal in enclosure.animals:
                return enclosure
        return None

    # ---------- Database ----------

    def load_data(self):
        """Reads the Excel file and creates Python objects for all data."""
        if not os.path.exists(self.db_file):
            print(f"No existing database found at {self.db_file}. Starting fresh.")
            return

        wb = openpyxl.load_workbook(self.db_file)

        # 1. Load Animals
        if "Animals" in wb.sheetnames:
            ws_animals = wb["Animals"]
            # min_row=2 skips the header row
            for row in ws_animals.iter_rows(min_row=2, values_only=True):
                if row[0] is None:  # Skip empty rows
                    continue
                # Older databases have no "Is Hungry" / "Diet" columns
                is_hungry = bool(row[7]) if len(row) > 7 and row[7] is not None else True
                diet = row[8] if len(row) > 8 and row[8] is not None else "Omnivore"
                animal = create_animal(diet, id=row[0], name=row[1], species=row[2], kingdom=row[3],
                                       age=row[4], health=row[5], habitat=row[6], is_hungry=is_hungry)
                self.animals.append(animal)

        # 2. Load Staff
        if "Staff" in wb.sheetnames:
            ws_staff = wb["Staff"]
            for row in ws_staff.iter_rows(min_row=2, values_only=True):
                if row[0] is None:
                    continue
                emp_id, name, role = row[0], row[1], row[2]
                salary = row[3] if len(row) > 3 and row[3] is not None else DEFAULT_SALARIES.get(role)
                # Create the correct subclass based on the role column
                staff_class = STAFF_CLASSES.get(role)
                if staff_class:
                    self.staff.append(staff_class(name=name, employee_id=emp_id, salary=salary))

        # 3. Load Enclosures (and re-link the animals living in them)
        if "Enclosures" in wb.sheetnames:
            ws_enclosures = wb["Enclosures"]
            for row in ws_enclosures.iter_rows(min_row=2, values_only=True):
                if row[0] is None:
                    continue
                enclosure = Enclosure(name=row[0], biome=row[1], capacity=row[2])
                occupant_ids = row[3] if len(row) > 3 and row[3] is not None else ""
                for animal_id in str(occupant_ids).split(","):
                    animal = self.find_animal(animal_id) if animal_id.strip() else None
                    if animal:
                        enclosure.animals.append(animal)
                self.enclosures.append(enclosure)

        print("Database loaded successfully.")

    def save_data(self):
        """Clears the Excel file and rewrites it with the current Python data."""
        wb = openpyxl.Workbook()

        # 1. Save Animals
        ws_animals = wb.active
        ws_animals.title = "Animals"
        ws_animals.append(["ID", "Name", "Species", "Kingdom", "Age", "Health", "Habitat", "Is Hungry", "Diet"])

        for a in self.animals:
            ws_animals.append([a.id, a.name, a.species, a.kingdom, a.age, a.health, a.habitat, a.is_hungry, a.diet])

        # 2. Save Staff
        ws_staff = wb.create_sheet(title="Staff")
        ws_staff.append(["Employee ID", "Name", "Role", "Salary"])

        for s in self.staff:
            ws_staff.append([s.employee_id, s.name, s.role, s.salary])

        # 3. Save Enclosures
        ws_enclosures = wb.create_sheet(title="Enclosures")
        ws_enclosures.append(["Enclosure Name", "Biome", "Max Capacity", "Occupant IDs"])

        for e in self.enclosures:
            occupant_ids = ",".join(a.id for a in e.animals)
            ws_enclosures.append([e.name, e.biome, e.capacity, occupant_ids])

        # Save and overwrite the file
        wb.save(self.db_file)
        print("Progress saved to database.")

    # ---------- Actions ----------

    def add_animal(self, id, name, species, kingdom, age, health, habitat, diet="Omnivore"):
        """Creates a new animal and adds it to the system."""
        if not str(id).strip():
            return "Error: Animal ID cannot be empty."
        if self.find_animal(id):
            return f"Error: An animal with ID '{id}' already exists."
        if not 0 <= health <= 100:
            return "Error: Health must be between 0 and 100."
        if age < 0:
            return "Error: Age cannot be negative."

        new_animal = create_animal(diet, str(id).strip(), name, species, kingdom, age, health, habitat)
        self.animals.append(new_animal)
        return f"Successfully registered {name}."

    def add_enclosure(self, name, biome, capacity):
        """Creates a new enclosure."""
        if not name.strip():
            return "Error: Enclosure name cannot be empty."
        if self.find_enclosure(name):
            return f"Error: An enclosure named '{name}' already exists."
        if capacity < 1:
            return "Error: Capacity must be at least 1."

        new_enclosure = Enclosure(name.strip(), biome, capacity)
        self.enclosures.append(new_enclosure)
        return f"Built new enclosure: {name}."

    def assign_animal_to_enclosure(self, animal_id, enclosure_name):
        """Finds an animal and an enclosure, then links them (moving it if needed)."""
        target_animal = self.find_animal(animal_id)
        if not target_animal:
            return f"Error: Could not find animal with ID '{animal_id}'."

        target_enclosure = self.find_enclosure(enclosure_name)
        if not target_enclosure:
            return f"Error: Could not find enclosure named '{enclosure_name}'."

        current_enclosure = self.find_enclosure_of(target_animal)
        if current_enclosure is target_enclosure:
            return f"{target_animal.name} is already in {target_enclosure.name}."

        result = target_enclosure.add_animal(target_animal)
        # Only remove from the old enclosure once the move has succeeded
        if result.startswith("Success") and current_enclosure:
            current_enclosure.remove_animal(target_animal.id)
            result += f" (moved from {current_enclosure.name})"
        return result

    def feed_animal(self, animal_id, food_type):
        animal = self.find_animal(animal_id)
        if not animal:
            return f"Error: Could not find animal with ID '{animal_id}'."
        return animal.feeding(food_type)

    def hire_staff(self, employee_id, name, role, salary=None):
        if not str(employee_id).strip():
            return "Error: Employee ID cannot be empty."
        if self.find_staff(employee_id):
            return f"Error: An employee with ID '{employee_id}' already exists."
        staff_class = STAFF_CLASSES.get(role)
        if not staff_class:
            return f"Error: Unknown role '{role}'."
        if salary is None:
            salary = DEFAULT_SALARIES[role]
        if salary < 0:
            return "Error: Salary cannot be negative."

        self.staff.append(staff_class(name=name, employee_id=str(employee_id).strip(), salary=salary))
        return f"Hired {name} as {role}."

    def treat_animal(self, vet_id, animal_id):
        vet = self.find_staff(vet_id)
        if not isinstance(vet, Vet):
            return f"Error: No veterinarian with ID '{vet_id}'."
        animal = self.find_animal(animal_id)
        if not animal:
            return f"Error: Could not find animal with ID '{animal_id}'."
        return vet.treat_animal(animal)
